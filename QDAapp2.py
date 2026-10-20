"""
QDA Classifier - Streamlit web app
=====================================
Upload a "database" Excel file (each sheet = one class, rows = samples,
columns = element concentrations) to train a Quadratic Discriminant
Analysis model, then upload an "unknown samples" Excel file to classify
each row.

Run with:
    streamlit run qda_app.py

Requires: streamlit, pandas, numpy, scikit-learn, plotly, openpyxl
Install with:
    pip install streamlit pandas numpy scikit-learn plotly openpyxl
"""

import io
import itertools

import numpy as np
import pandas as pd
import plotly.graph_objects as go
import streamlit as st
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from sklearn.decomposition import PCA
from sklearn.discriminant_analysis import QuadraticDiscriminantAnalysis
from sklearn.metrics import accuracy_score, classification_report, confusion_matrix
from sklearn.model_selection import train_test_split
from sklearn.preprocessing import StandardScaler

# Fixed colors for the known glass classes; anything else falls back to the pool below.
FIXED_CLASS_COLORS = {
    "PED glass": "green",
    "Container glass": "red",
    "Float glass": "blue",
}
FALLBACK_COLOR_POOL = ["purple", "orange", "brown", "magenta", "gray", "teal", "gold"]

# Fixed model setting (no longer a slider).
REG_PARAM = 0.0  # 0 = full QDA; raise toward 1 if classes are small vs. number of elements

st.set_page_config(page_title="QDA Classifier", layout="wide")
st.title("QDA Classifier")

st.markdown(
    """
Upload a **database file** where each sheet is a class (e.g. `Float glass`,
`Container glass`, `PED glass`) and rows are samples with element
concentrations as columns. Then upload an **unknown samples file** with
the same element columns to classify.
"""
)

# ---------------------------------------------------------------
# Session state so results survive re-runs / interactions
# ---------------------------------------------------------------
if "model" not in st.session_state:
    st.session_state.model = None
    st.session_state.scaler = None
    st.session_state.feature_cols = None
    st.session_state.class_labels = None
    st.session_state.log_transform = None
    st.session_state.random_state = None
    st.session_state.pca = None
    st.session_state.viz_qda = None
    st.session_state.X_pca = None
    st.session_state.X_ref = None
    st.session_state.y_ref = None


def _class_color_map(classes):
    pool = itertools.cycle(FALLBACK_COLOR_POOL)
    color_map = {}
    for c in classes:
        if c in FIXED_CLASS_COLORS:
            color_map[c] = FIXED_CLASS_COLORS[c]
        else:
            color_map[c] = next(pool)
    return color_map


def plot_class_regions(pca, viz_qda, X_pca, y, unknown_pca=None, unknown_labels=None,
                        unknown_ids=None, highlight_idx=None, grid_res=180):
    classes = sorted(y.unique())
    n = len(classes)
    color_map = _class_color_map(classes)

    pad_x = 0.08 * ((X_pca[:, 0].max() - X_pca[:, 0].min()) or 1)
    pad_y = 0.08 * ((X_pca[:, 1].max() - X_pca[:, 1].min()) or 1)
    x_min, x_max = X_pca[:, 0].min() - pad_x, X_pca[:, 0].max() + pad_x
    y_min, y_max = X_pca[:, 1].min() - pad_y, X_pca[:, 1].max() + pad_y
    xx, yy = np.meshgrid(np.linspace(x_min, x_max, grid_res), np.linspace(y_min, y_max, grid_res))
    class_index = {c: i for i, c in enumerate(classes)}
    Z_labels = viz_qda.predict(np.c_[xx.ravel(), yy.ravel()])
    Zi = np.array([class_index[c] for c in Z_labels]).reshape(xx.shape)

    colorscale = []
    for i, c in enumerate(classes):
        colorscale.append([i / n, color_map[c]])
        colorscale.append([(i + 1) / n, color_map[c]])

    fig = go.Figure()
    fig.add_trace(go.Heatmap(
        x=xx[0], y=yy[:, 0], z=Zi,
        colorscale=colorscale, zmin=0, zmax=n - 1,
        showscale=False, opacity=0.22, hoverinfo="skip",
    ))
    for c in classes:
        mask = (y == c).values
        fig.add_trace(go.Scatter(
            x=X_pca[mask, 0], y=X_pca[mask, 1], mode="markers",
            marker=dict(color=color_map[c], size=8, line=dict(color="black", width=0.5)),
            name=f"{c} (n={mask.sum()})",
            hovertemplate=f"Class: {c}<br>PC1: %{{x:.2f}} | PC2: %{{y:.2f}}<extra></extra>",
        ))
    if unknown_pca is not None and len(unknown_pca):
        labels = list(unknown_labels) if unknown_labels is not None else [""] * len(unknown_pca)
        ids = list(unknown_ids) if unknown_ids is not None else [str(i + 1) for i in range(len(unknown_pca))]
        customdata = list(zip(ids, labels))
        fig.add_trace(go.Scatter(
            x=unknown_pca[:, 0], y=unknown_pca[:, 1], mode="markers",
            marker=dict(symbol="circle", size=10, color="white", line=dict(width=2, color="black")),
            customdata=customdata, name="Unknown sample(s)",
            hovertemplate="Sample: %{customdata[0]}<br>Predicted: %{customdata[1]}"
                           "<br>PC1: %{x:.2f} | PC2: %{y:.2f}<extra></extra>",
        ))
        if highlight_idx is not None and 0 <= highlight_idx < len(unknown_pca):
            fig.add_trace(go.Scatter(
                x=[unknown_pca[highlight_idx, 0]], y=[unknown_pca[highlight_idx, 1]],
                mode="markers",
                marker=dict(symbol="circle-open", size=26, color="gold", line=dict(width=4, color="gold")),
                name="Selected sample", hoverinfo="skip", showlegend=False,
            ))
    fig.update_layout(
        title="QDA projection (shaded decision regions)",
        xaxis_title=f"PC1 ({pca.explained_variance_ratio_[0]:.1%} of variance)",
        yaxis_title=f"PC2 ({pca.explained_variance_ratio_[1]:.1%} of variance)",
        legend=dict(orientation="h", yanchor="bottom", y=1.02, x=0),
        height=550, margin=dict(t=80),
    )
    return fig


def qda_pairwise_boundary_coeffs(qda, class_i, class_j):
    """Coefficients (a, b, c, d, e, f) of the quadratic boundary curve

        a*x1^2 + b*x1*x2 + c*x2^2 + d*x1 + e*x2 + f = 0

    between two classes of a 2-D QuadraticDiscriminantAnalysis fitted with
    store_covariance=True. This is the curve where the two classes'
    discriminant functions are equal, i.e. delta_i(x) = delta_j(x), which is
    exactly the boundary drawn between two shaded regions in the projection
    plot.
    """
    classes = list(qda.classes_)
    i, j = classes.index(class_i), classes.index(class_j)
    mu_i, mu_j = np.asarray(qda.means_[i]), np.asarray(qda.means_[j])
    Sigma_i, Sigma_j = np.asarray(qda.covariance_[i]), np.asarray(qda.covariance_[j])
    pi_i, pi_j = qda.priors_[i], qda.priors_[j]

    Mi = np.linalg.inv(Sigma_i)
    Mj = np.linalg.inv(Sigma_j)
    dM = Mi - Mj  # symmetric 2x2

    a = -0.5 * dM[0, 0]
    c = -0.5 * dM[1, 1]
    b = -dM[0, 1]  # dM is symmetric, so this covers the cross term

    lin = mu_i @ Mi - mu_j @ Mj
    d, e = lin[0], lin[1]

    f = (
        -0.5 * (mu_i @ Mi @ mu_i - mu_j @ Mj @ mu_j)
        - 0.5 * (np.log(np.linalg.det(Sigma_i)) - np.log(np.linalg.det(Sigma_j)))
        + np.log(pi_i / pi_j)
    )
    return a, b, c, d, e, f


def format_quadratic_latex(a, b, c, d, e, f, x1="x_1", x2="x_2"):
    """Render a*x1^2 + b*x1*x2 + c*x2^2 + d*x1 + e*x2 + f = 0 as a LaTeX string."""
    def term(coef, suffix):
        if abs(coef) < 1e-10:
            return None
        return f"{coef:.4f}{suffix}"

    parts = [
        term(a, f"{x1}^2"),
        term(b, f"{x1}{x2}"),
        term(c, f"{x2}^2"),
        term(d, f"{x1}"),
        term(e, f"{x2}"),
        term(f, ""),
    ]
    parts = [p for p in parts if p is not None]
    if not parts:
        return "0 = 0"

    expr = parts[0]
    for p in parts[1:]:
        if p.startswith("-"):
            expr += f" - {p[1:]}"
        else:
            expr += f" + {p}"
    return expr + " = 0"


# ---------------------------------------------------------------
# Excel export
# ---------------------------------------------------------------
def _py(v):
    """Convert numpy scalars / NaN to plain Python values openpyxl can write."""
    if v is None:
        return None
    try:
        if pd.isna(v):
            return None
    except (TypeError, ValueError):
        pass
    return v.item() if hasattr(v, "item") else v


def _write_header(ws, row, headers):
    for j, h in enumerate(headers, start=1):
        cell = ws.cell(row=row, column=j, value=h)
        cell.font = Font(name="Calibri", size=11, bold=True)
        cell.alignment = Alignment(horizontal="center")


def _set_widths(ws, headers, first_col_width=None):
    for j, h in enumerate(headers, start=1):
        ws.column_dimensions[get_column_letter(j)].width = max(11, len(str(h)) + 3)
    if first_col_width:
        ws.column_dimensions["A"].width = first_col_width


def qda_class_score_terms(model, scaler):
    """Per-class pieces of the full quadratic score, expressed in the measurement
    space the user supplies (log10 concentrations if the log transform is on),
    so no standardization step is needed:

        score_k(x) = -0.5 (x - mean_k)^T Precision_k (x - mean_k) + Constant_k

    Constant_k = -0.5 ln|Sigma_k| + ln(prior_k). The posterior probabilities are
    the softmax of these scores.
    """
    D = np.asarray(scaler.scale_)
    out = {}
    for k, cls in enumerate(model.classes_):
        S = np.asarray(model.covariance_[k])          # covariance in standardized space
        mean_x = np.asarray(scaler.mean_) + D * np.asarray(model.means_[k])
        precision_x = np.linalg.inv(S) / np.outer(D, D)
        logdet_x = np.linalg.slogdet(S)[1] + 2.0 * np.sum(np.log(D))
        const = -0.5 * logdet_x + np.log(model.priors_[k])
        out[cls] = (const, mean_x, precision_x)
    return out


def pc_axis_coeffs(pca, scaler):
    """Intercept and per-feature coefficients so that PC_k = intercept_k + sum(coef * feature)."""
    scale = np.asarray(scaler.scale_)
    mean = np.asarray(scaler.mean_)
    comp = np.asarray(pca.components_)
    coefs = comp / scale
    intercepts = -(comp @ (mean / scale + np.asarray(pca.mean_)))
    return intercepts, coefs


def build_excel(results, unk_pca, model, scaler, pca, viz_qda, feature_cols,
                log_transform, prob_threshold):
    """Return the classification results as .xlsx bytes (same layout as the LDA export):
    Predictions + the formula tabs describing the fitted model."""
    wb = Workbook()
    space = "log10 of the concentration" if log_transform else "the raw concentration"
    classes = list(model.classes_)

    # ---- Predictions ----------------------------------------------------------
    ws = wb.active
    ws.title = "Predictions"
    df = results.copy()
    df["PC1"] = unk_pca[:, 0]
    df["PC2"] = unk_pca[:, 1]
    headers = [str(c) for c in df.columns]

    ws["A1"] = (
        "Each row is an unknown sample with its predicted class, posterior probability, "
        f"confidence flag (threshold = {prob_threshold:.0%}), per-class probabilities, "
        "and its position on the PC1/PC2 projection plot."
    )
    ws["A2"] = (
        "PC1 and PC2 are computed as axis = Intercept + sum(coefficient x feature value) — see the "
        "'PC1_PC2 formula' tab. The class itself is decided in the full feature space by the "
        "quadratic scores on the 'Classification formula' tab. Feature value = " + space + "."
    )
    _write_header(ws, 4, headers)
    conf_col = headers.index("confident") if "confident" in headers else None
    amber = PatternFill("solid", start_color="FFF3CD", end_color="FFF3CD")
    for i, row in enumerate(df.itertuples(index=False), start=5):
        for j, v in enumerate(row, start=1):
            ws.cell(row=i, column=j, value=_py(v))
        if conf_col is not None and not bool(row[conf_col]):
            for j in range(1, len(headers) + 1):
                ws.cell(row=i, column=j).fill = amber
    _set_widths(ws, headers)
    ws.column_dimensions["A"].width = max(14, int(df.iloc[:, 0].astype(str).str.len().max() or 0) + 3)

    # ---- Classification formula ----------------------------------------------
    ws = wb.create_sheet("Classification formula")
    ws["A1"] = ("QDA assigns each class a quadratic discriminant score. A new sample is classified "
                "into the class with the highest score:")
    ws["A2"] = ("score(class) = Constant - 0.5 x (x - Mean)ᵀ x Precision x (x - Mean),   "
                "where x = vector of feature values (" + space + ")")
    ws["A3"] = ("Constant = -0.5 x ln|Covariance| + ln(prior). Precision = inverse covariance matrix "
                "(one block per class below). Posterior probabilities = softmax of the scores.")
    ws["A4"] = "Class means (one row per class) are listed in the feature columns of the first table."
    terms = qda_class_score_terms(model, scaler)
    head = ["Class", "Constant"] + list(feature_cols)
    _write_header(ws, 6, head)
    for i, cls in enumerate(classes, start=7):
        const, mean_x, _ = terms[cls]
        c = ws.cell(row=i, column=1, value=cls)
        c.font = Font(name="Calibri", size=11, bold=True)
        ws.cell(row=i, column=2, value=float(const))
        for j, m in enumerate(mean_x, start=3):
            ws.cell(row=i, column=j, value=float(m))
    r = 7 + len(classes) + 1
    for cls in classes:
        _, _, P = terms[cls]
        t = ws.cell(row=r, column=1, value=f"{cls} - Precision matrix (inverse covariance)")
        t.font = Font(name="Calibri", size=11, bold=True)
        _write_header(ws, r + 1, [""] + list(feature_cols))
        for a, feat in enumerate(feature_cols):
            c = ws.cell(row=r + 2 + a, column=1, value=feat)
            c.font = Font(name="Calibri", size=11, bold=True)
            for b in range(len(feature_cols)):
                ws.cell(row=r + 2 + a, column=2 + b, value=float(P[a, b]))
        r += len(feature_cols) + 4
    _set_widths(ws, head, first_col_width=34)

    # ---- Pairwise boundaries --------------------------------------------------
    ws = wb.create_sheet("Pairwise boundaries")
    ws["A1"] = "The boundary between two classes is where their scores are tied. In the PC1/PC2 projection plot it is the curve:"
    ws["A2"] = "a x PC1² + b x PC1·PC2 + c x PC2² + d x PC1 + e x PC2 + f = 0"
    ws["A3"] = ("Positive favors the first-named class in the 'Boundary' column, negative favors the second. "
                "These describe the 2-component visualization only, not the full-space classifier.")
    head = ["Boundary", "a (PC1²)", "b (PC1·PC2)", "c (PC2²)", "d (PC1)", "e (PC2)", "f (constant)"]
    _write_header(ws, 5, head)
    for i, (ci, cj) in enumerate(itertools.combinations(sorted(classes), 2), start=6):
        c = ws.cell(row=i, column=1, value=f"{ci} vs {cj}")
        c.font = Font(name="Calibri", size=11, bold=True)
        for j, v in enumerate(qda_pairwise_boundary_coeffs(viz_qda, ci, cj), start=2):
            ws.cell(row=i, column=j, value=float(v))
    _set_widths(ws, head, first_col_width=34)

    # ---- PC1_PC2 formula ------------------------------------------------------
    ws = wb.create_sheet("PC1_PC2 formula")
    ws["A1"] = ("This is the formula behind the projection axes used in the app's plots: each axis is a "
                "linear combination of your feature values (" + space + ").")
    ws["A2"] = "axis = Intercept + sum(coefficient x feature value)"
    head = ["Component", "Intercept"] + list(feature_cols)
    _write_header(ws, 4, head)
    intercepts, coefs = pc_axis_coeffs(pca, scaler)
    for k in range(2):
        c = ws.cell(row=5 + k, column=1, value=f"PC{k + 1}")
        c.font = Font(name="Calibri", size=11, bold=True)
        ws.cell(row=5 + k, column=2, value=float(intercepts[k]))
        for j, v in enumerate(coefs[k], start=3):
            ws.cell(row=5 + k, column=j, value=float(v))
    _set_widths(ws, head, first_col_width=34)

    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


# =================================================================
# STEP 1 — Upload database & train
# =================================================================
st.header("1. Train on your database")

db_file = st.file_uploader("Database Excel file (.xlsx)", type=["xlsx"], key="db")

if db_file is not None:
    try:
        xls = pd.ExcelFile(db_file)
    except Exception as e:
        st.error(f"Could not read Excel file: {e}")
        st.stop()

    sheet_names = xls.sheet_names
    st.write(f"Found {len(sheet_names)} sheets (= classes): {', '.join(sheet_names)}")

    selected_sheets = st.multiselect("Sheets to use as classes", sheet_names, default=sheet_names)

    if len(selected_sheets) < 2:
        st.warning("Select at least 2 sheets/classes to run QDA.")
    else:
        preview_cols = list(pd.read_excel(xls, sheet_name=selected_sheets[0], nrows=0).columns)
        id_col_choice = st.selectbox(
            "Sample ID column to exclude from training (not a measurement)",
            options=["(none)"] + preview_cols,
            index=1 if preview_cols else 0,
            help="Pick the column that identifies each sample (e.g. Sample_ID), "
                 "so it isn't treated as an element measurement. Defaults to the first column.",
        )

        log_transform = st.checkbox(
            "Log10-transform concentrations", value=True,
            help="Recommended: trace-element concentrations are typically right-skewed; "
                 "QDA assumes each class is multivariate normal.",
        )
        test_size = st.slider(
            "Test size", 0.1, 0.5, 0.2, 0.05,
            help="Fraction of the database held out to evaluate the model. "
                 "The model is trained on the rest, then scored on the held-out samples.",
        )
        random_state = st.number_input(
            "Random state", min_value=0, max_value=10_000, value=42, step=1,
            help="Seed used for the train/test split and the PCA visualization. "
                 "Change it to check how sensitive the results are, or keep it fixed for reproducibility.",
        )

        if st.button("Train model", type="primary"):
            frames = []
            for sheet in selected_sheets:
                sheet_df = pd.read_excel(xls, sheet_name=sheet)
                sheet_df = sheet_df.dropna(how="all")
                sheet_df["label"] = sheet
                frames.append(sheet_df)
            data = pd.concat(frames, ignore_index=True)

            excluded_cols = {"label"}
            if id_col_choice != "(none)":
                excluded_cols.add(id_col_choice)
            feature_cols = [c for c in data.columns if c not in excluded_cols]

            X = data[feature_cols].apply(pd.to_numeric, errors="coerce")
            y = data["label"]

            mask = X.notna().all(axis=1) & y.notna()
            n_dropped = (~mask).sum()
            if n_dropped:
                st.warning(f"Dropped {n_dropped} row(s) with missing/non-numeric values.")
            X, y = X[mask], y[mask]

            dropped_log = []
            if log_transform:
                X = np.log10(X.replace(0, np.nan))
                keep = X.dropna().index
                dropped_log = sorted(set(X.index) - set(keep))
                X, y = X.loc[keep], y.loc[keep]
                if dropped_log:
                    st.warning(
                        f"Dropped {len(dropped_log)} more row(s) with zero/invalid values "
                        f"before the log10 transform."
                    )

            st.write(f"Using {len(feature_cols)} feature columns: {', '.join(feature_cols)}")
            st.write("Class counts:")
            st.dataframe(y.value_counts().rename("samples"))

            labels = sorted(y.unique())

            # Each class needs at least 2 samples for a stratified split
            if y.value_counts().min() < 2:
                st.error("Every class needs at least 2 samples for a train/test split.")
                st.stop()

            n_test = int(np.ceil(test_size * len(y)))
            if n_test < len(labels):
                st.error("Test set is smaller than the number of classes. Increase the test size.")
                st.stop()

            X_train, X_test, y_train, y_test = train_test_split(
                X, y, test_size=test_size, stratify=y, random_state=random_state
            )

            # Evaluation model: scaler and QDA fitted on the training portion only,
            # so the test score isn't inflated by information from the test samples.
            eval_scaler = StandardScaler().fit(X_train)
            eval_qda = QuadraticDiscriminantAnalysis(reg_param=REG_PARAM)
            eval_qda.fit(eval_scaler.transform(X_train), y_train)
            y_pred_test = eval_qda.predict(eval_scaler.transform(X_test))

            col1, col2 = st.columns(2)
            with col1:
                st.metric("Test-set accuracy", f"{accuracy_score(y_test, y_pred_test):.1%}")
                st.caption(
                    f"Trained on {len(y_train)} samples, tested on {len(y_test)} "
                    f"({test_size:.0%} held out, stratified)."
                )
                st.text("Classification report:")
                st.text(classification_report(
                    y_test, y_pred_test, labels=labels, target_names=labels, zero_division=0
                ))
            with col2:
                cm = confusion_matrix(y_test, y_pred_test, labels=labels)
                st.text("Confusion matrix (rows = true, columns = predicted):")
                st.dataframe(pd.DataFrame(cm, index=labels, columns=labels))

            # Final model trained on ALL data (used for real classification)
            scaler = StandardScaler().fit(X)
            Xs = scaler.transform(X)
            # store_covariance=True so the Excel export can include the exact class
            # covariance matrices (and thus the full quadratic scoring formula).
            final_qda = QuadraticDiscriminantAnalysis(reg_param=REG_PARAM, store_covariance=True)
            final_qda.fit(Xs, y)

            st.session_state.model = final_qda
            st.session_state.scaler = scaler
            st.session_state.feature_cols = feature_cols
            st.session_state.class_labels = labels
            st.session_state.log_transform = log_transform
            st.session_state.random_state = random_state

            st.success(
                "Final model trained on all available data. "
                "You can now classify unknown samples below."
            )

            # ---------------------------------------------------
            # Example class-region plot
            # ---------------------------------------------------
            st.subheader("QDA projection")
            pca = PCA(n_components=2, random_state=random_state).fit(Xs)
            X_pca = pca.transform(Xs)
            # store_covariance=True so we can pull each class's fitted covariance
            # matrix back out afterwards to build the boundary formulas below.
            viz_qda = QuadraticDiscriminantAnalysis(
                reg_param=REG_PARAM, store_covariance=True
            ).fit(X_pca, y)

            st.session_state.pca = pca
            st.session_state.viz_qda = viz_qda
            st.session_state.X_pca = X_pca
            st.session_state.X_ref = X
            st.session_state.y_ref = y

            plot_tab, formula_tab = st.tabs(["Projection plot", "Decision boundary formulas"])

            with plot_tab:
                st.caption(
                    "The real model uses all features at once, which can't be drawn directly. "
                    "This plot compresses the data to its 2 strongest principal components "
                    "and shows the region each class would occupy — an approximation of the "
                    "true decision boundary, useful for building intuition."
                )
                st.plotly_chart(plot_class_regions(pca, viz_qda, X_pca, y), use_container_width=True)

            with formula_tab:
                st.caption(
                    "For each pair of classes, QDA assigns a point to whichever class has the "
                    "higher discriminant score. Both scores are quadratic functions of the "
                    "coordinates, so the boundary you see between two shaded regions above is "
                    r"the curve where the two scores are equal — solving $\delta_i(x)=\delta_j(x)$ "
                    "gives a quadratic equation in "
                    r"$x_1$ (PC1) and $x_2$ (PC2):"
                )
                st.latex(r"a\,x_1^2 + b\,x_1 x_2 + c\,x_2^2 + d\,x_1 + e\,x_2 + f = 0")
                st.caption(
                    "Below are the fitted coefficients for every pair of classes, in the same "
                    "PC1/PC2 space as the plot above. Note these describe only the 2-component "
                    "visualization — the actual classifier used in step 2 works in the full "
                    f"{len(feature_cols)}-element space and can't be reduced to a single formula "
                    "this simply."
                )
                for ci, cj in itertools.combinations(sorted(y.unique()), 2):
                    a, b, c, d, e, f_coef = qda_pairwise_boundary_coeffs(viz_qda, ci, cj)
                    st.markdown(f"**{ci} vs. {cj}**")
                    st.latex(format_quadratic_latex(a, b, c, d, e, f_coef))


# =================================================================
# STEP 2 — Upload unknown samples & classify
# =================================================================
st.header("2. Classify unknown samples")

if st.session_state.model is None:
    st.info("Train a model in step 1 first.")
else:
    feature_cols = st.session_state.feature_cols
    scaler = st.session_state.scaler
    model = st.session_state.model
    log_transform = st.session_state.log_transform

    prob_threshold = st.slider(
        "Confidence threshold", 0.5, 1.0, 0.90, 0.01,
        help="Posterior probability below which a prediction is flagged as not confident.",
    )

    tab_upload, tab_manual = st.tabs(["Upload a file", "Enter one sample manually"])

    def _classify(unknown_df):
        missing_cols = [c for c in feature_cols if c not in unknown_df.columns]
        if missing_cols:
            st.error(f"Unknown samples file is missing required columns: {missing_cols}")
            return None

        X_unk = unknown_df[feature_cols].apply(pd.to_numeric, errors="coerce")
        if log_transform:
            X_unk = np.log10(X_unk.replace(0, np.nan))

        valid_mask = X_unk.notna().all(axis=1)
        n_invalid = (~valid_mask).sum()
        if n_invalid:
            st.warning(f"Skipping {n_invalid} row(s) with missing/invalid feature values.")

        X_unk_valid = X_unk[valid_mask]
        if len(X_unk_valid) == 0:
            st.error("No valid rows to classify after cleaning.")
            return None

        Xs_unk = scaler.transform(X_unk_valid)
        pred = model.predict(Xs_unk)
        proba = model.predict_proba(Xs_unk)
        proba_df = pd.DataFrame(proba, index=X_unk_valid.index, columns=model.classes_)

        results = unknown_df.loc[valid_mask].copy()
        results["Predicted_class"] = pred
        results["posterior_probability"] = proba_df.max(axis=1).values
        results["confident"] = results["posterior_probability"] >= prob_threshold
        for cls in model.classes_:
            results[f"P({cls})"] = proba_df[cls].values
        return results

    with tab_upload:
        unknown_file = st.file_uploader(
            "Unknown samples Excel/CSV file (.xlsx or .csv)", type=["xlsx", "csv"], key="unknown"
        )
        if unknown_file is not None:
            if unknown_file.name.lower().endswith(".csv"):
                unknown_df = pd.read_csv(unknown_file)
            else:
                unknown_df = pd.read_excel(unknown_file)
            unknown_df = unknown_df.dropna(how="all")

            results = _classify(unknown_df)
            if results is not None:
                def _highlight(row):
                    return ["background-color: #fff3cd" if not row["confident"] else "" for _ in row]

                st.subheader("Results")
                st.caption("Click a row to highlight that sample in the projection below.")
                selection_event = st.dataframe(
                    results.style.apply(_highlight, axis=1).format(
                        {"posterior_probability": "{:.4f}"}
                        | {c: "{:.4f}" for c in results.columns if c.startswith("P(")}
                    ),
                    use_container_width=True,
                    on_select="rerun",
                    selection_mode="single-row",
                    key="results_table",
                )
                selected_rows = (
                    selection_event.selection.rows
                    if selection_event is not None and selection_event.selection is not None
                    else []
                )
                highlight_idx = selected_rows[0] if selected_rows else None

                n_low = (~results["confident"]).sum()
                if n_low:
                    st.warning(f"{n_low} sample(s) fell below the {prob_threshold:.0%} confidence threshold.")

                # Build the PCA points from the SAME rows/order as `results` (i.e. only the
                # valid, classified rows) so the table's row selection lines up correctly
                # with the corresponding point below.
                X_unk_for_pca = unknown_df.loc[results.index, feature_cols].apply(pd.to_numeric, errors="coerce")
                if log_transform:
                    X_unk_for_pca = np.log10(X_unk_for_pca.replace(0, np.nan))
                unk_pca = st.session_state.pca.transform(scaler.transform(X_unk_for_pca))
                # Use a non-feature column (e.g. Sample_ID) to label unknown points on hover,
                # falling back to a simple row number if no such column exists.
                id_cols = [c for c in unknown_df.columns if c not in feature_cols]
                sample_ids = results[id_cols[0]] if id_cols else (results.index + 1)

                st.plotly_chart(
                    plot_class_regions(
                        st.session_state.pca, st.session_state.viz_qda, st.session_state.X_pca,
                        st.session_state.y_ref, unknown_pca=unk_pca,
                        unknown_labels=results["Predicted_class"],
                        unknown_ids=sample_ids,
                        highlight_idx=highlight_idx,
                    ),
                    use_container_width=True,
                )

                dl_xlsx, dl_csv = st.columns(2)
                with dl_xlsx:
                    st.download_button(
                        "Download results as Excel",
                        data=build_excel(
                            results, unk_pca, model, scaler, st.session_state.pca,
                            st.session_state.viz_qda, feature_cols, log_transform, prob_threshold,
                        ),
                        file_name="classification_results.xlsx",
                        mime="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
                        type="primary",
                    )
                with dl_csv:
                    st.download_button(
                        "Download results as CSV",
                        data=results.to_csv(index=False).encode("utf-8"),
                        file_name="classification_results.csv",
                        mime="text/csv",
                    )

    with tab_manual:
        st.caption("Enter concentrations for a single sample (same units as the reference database).")
        cols = st.columns(4)
        manual_values = {}
        for i, el in enumerate(feature_cols):
            with cols[i % 4]:
                manual_values[el] = st.number_input(el, value=0.0, format="%.6f", key=f"manual_{el}")

        if st.button("Classify this sample"):
            if all(v == 0.0 for v in manual_values.values()):
                st.warning("Enter at least some non-zero concentrations first.")
            else:
                manual_df = pd.DataFrame([manual_values])
                results = _classify(manual_df)
                if results is not None:
                    st.dataframe(results, use_container_width=True)
                    pred = results.loc[0, "Predicted_class"]
                    prob = results.loc[0, "posterior_probability"]
                    st.success(f"Predicted class: **{pred}** (posterior probability: {prob:.1%})")